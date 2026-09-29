import io
import unittest
from pptx import Presentation
from report_generator import AnalysisAgent, SlidePlannerAgent, PptxGeneratorAgent, percent, selected_periods


class PeriodTests(unittest.TestCase):
    def setUp(self):
        self.catalog=[dict(id=p,name=p,program=p,group='PHCs') for p in ('hh','ipccc','phc_ica')]
        self.entries=[dict(program='hh',facilityId='hh',period=f'M{m:02}',score=m*5) for m in range(1,13)]
        self.entries += [dict(program='ipccc',facilityId='ipccc',period=f'Q{q}',score=q*20) for q in range(1,5)]
        self.entries += [dict(program='phc_ica',facilityId='phc_ica',period=p,score=s) for p,s in [('H1',40),('H2',80)]]
        self.preview=dict(entries=self.entries,unmatched=[],skippedZeros=0,duplicateValues=0)

    def report(self,scope,program):
        return AnalysisAgent().run(self.preview,self.catalog,scope,program)[0]

    def test_all_quarters(self):
        for q in range(1,5):
            self.assertEqual(self.report(f'Q{q}','hh')['aggregate'],(q*3-1)*5)
            self.assertEqual(self.report(f'Q{q}','ipccc')['aggregate'],q*20)

    def test_halves_and_annual(self):
        for scope,expected in [('H1',40),('H2',80),('year',60)]:
            self.assertEqual(self.report(scope,'phc_ica')['aggregate'],expected)
        self.assertEqual(self.report('year','ipccc')['aggregate'],50)
        self.assertEqual(self.report('H2','hh')['aggregate'],47.5)

    def test_custom_comparison(self):
        r=self.report('compare:Q4,Q1,Q3','ipccc')
        self.assertEqual(r['periods'],['Q1','Q3','Q4'])
        self.assertAlmostEqual(r['aggregate'],160/3)
        self.assertEqual(r['common_change'],60)

    def test_half_not_distributed(self):
        self.assertIsNone(self.report('Q1','phc_ica')['aggregate'])
        self.assertEqual(self.report('compare:Q1,Q3','phc_ica')['periods'],[])

    def test_missing_not_zero(self):
        self.entries.pop(0)
        r=self.report('H1','hh')
        self.assertIsNone(r['aggregate'])
        self.assertEqual(r['partial'],1)
        self.assertEqual(r['common_count'],0)

    def test_true_zero(self):
        self.entries[12]['score']=0
        self.assertEqual(self.report('Q1','ipccc')['aggregate'],0)
        self.assertEqual(percent('٠٪'),0)

    def test_invalid_scopes(self):
        for scope in ['compare:','compare:Q1,Q1','Q5']:
            with self.assertRaises(ValueError): selected_periods(scope)

    def test_editable_deck(self):
        analysis=AnalysisAgent().run(self.preview,self.catalog,'year')
        deck=Presentation(io.BytesIO(PptxGeneratorAgent().run(SlidePlannerAgent().run(analysis,self.preview,2026,'year'),2026,scope='year')))
        self.assertGreater(len(deck.slides),10)
        self.assertTrue(any(s.has_chart for slide in deck.slides for s in slide.shapes))
        self.assertTrue(any(s.has_table for slide in deck.slides for s in slide.shapes))
        self.assertFalse(any(c in ''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame) for c in '٠١٢٣٤٥٦٧٨٩'))


if __name__=='__main__': unittest.main()

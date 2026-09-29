"""Five deterministic agents: Tracker XLSX -> pandas -> editable quarterly, half-year and annual PPTX.

Adapted from infection-control-reports. No LLM or network calls. Runs both in
CPython and Pyodide; input percentages are never converted into fake counts.
"""
from __future__ import annotations

import argparse
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from python_engine.base_pipeline import PptxGeneratorAgent as BaseRenderer, ReportConfig, SheetSpec

LABELS = {"ipccc": "برنامج مكافحة العدوى IPCCC", "hh": "نظافة اليدين", "rpp": "الوقاية التنفسية RPP", "phc_ica": "التقييم الذاتي PHC ICA"}
GROUPS = {"Government MOH Hospitals": "المستشفيات الحكومية", "Private Hospitals": "المستشفيات الخاصة", "Non-MOH Government Hospital": "مستشفيات خارج وزارة الصحة", "Dental Centers": "مراكز الأسنان", "HD Private Centers": "الغسيل الكلوي الخاص", "HD Charity Centers": "الغسيل الكلوي الخيري", "PHCs": "الرعاية الصحية الأولية"}
MONTHS = {name: f"M{i:02}" for i, name in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
MAX_FILE_BYTES = 12 * 1024 * 1024


def normalize(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9\u0600-\u06ff]+", " ", str(value).translate(DIGITS).lower().replace("&", " and "))).strip()


def percent(value):
    if value is None or pd.isna(value) or str(value).strip() == "":
        return None
    raw = str(value).translate(DIGITS).strip()
    explicit = "%" in raw or "٪" in raw
    try:
        number = float(raw.replace("%", "").replace("٪", "").replace("٫", ".").replace(",", ".").replace(" ", ""))
    except (ValueError, TypeError):
        raise ValueError("قيمة نسبة غير رقمية في الملف.") from None
    if not explicit and 0 < number <= 1:
        number *= 100
    if not math.isfinite(number) or not 0 <= number <= 100:
        raise ValueError("النسب يجب أن تكون بين 0 و100.")
    return round(number, 2)


def fmt(value):
    return "غير مكتمل" if value is None or pd.isna(value) else f"{value:.1f}%"


def delta(value):
    return "—" if value is None or pd.isna(value) else f"{value:+.1f}"


def program_for(name):
    key = normalize(name)
    if "ipccc" in key:
        return "ipccc"
    if re.search(r"\brpp\b", key):
        return "rpp"
    if "phc" in key and ("ica" in key or "self assess" in key):
        return "phc_ica"
    if "h h" in key or "hand hygiene" in key:
        return "hh"
    return None


def period_for(value, program):
    key = normalize(value)
    if program == "hh":
        return MONTHS.get(key)
    if program == "phc_ica":
        return "H1" if re.search(r"1(st)? half", key) else "H2" if re.search(r"2(nd)? half", key) else None
    for i, ordinal in enumerate(["1st", "2nd", "3rd", "4th"], 1):
        if re.search(rf"\b({ordinal}|q{i}|quarter {i})\b", key) or (i == 4 and "4rd" in key):
            return f"Q{i}"
    return None


class DataExtractorAgent:
    """Explicit adapter for the five supplied wide Tracker sheets."""

    @staticmethod
    def inspect(source):
        book = pd.ExcelFile(io.BytesIO(source), engine="openpyxl")
        try:
            return {name: book.parse(name, header=None, nrows=3).fillna("").values.tolist() for name in book.sheet_names}
        finally:
            book.close()

    def run(self, source: bytes, catalog: list[dict], skip_zeros=True):
        if not source or len(source) > MAX_FILE_BYTES:
            raise ValueError("اختر ملف XLSX بحجم لا يتجاوز 12 ميغابايت.")
        try:
            with zipfile.ZipFile(io.BytesIO(source)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(x.file_size for x in infos) > 96 * 1024 * 1024:
                    raise ValueError("الملف أكبر من حدود المعالجة. قسّمه إلى ملفات أصغر.")
                if "xl/workbook.xml" not in archive.namelist():
                    raise ValueError("الملف ليس مصنف Excel صالحًا.")
        except zipfile.BadZipFile:
            raise ValueError("تعذر فتح ملف Excel. تأكد من صيغة XLSX.") from None
        lookup = {(f["program"], normalize(n)): f for f in catalog for n in [f["name"], f.get("arabicName", "")] if normalize(n)}
        ids = {f["id"]: f for f in catalog}
        book = pd.ExcelFile(io.BytesIO(source), engine="openpyxl")
        found, unmatched, sheets, seen = [], set(), [], {}
        zeros = duplicates = 0
        try:
            for name in book.sheet_names:
                program = program_for(name)
                if not program:
                    continue
                frame = book.parse(name, header=None, nrows=2001)
                if len(frame) > 2000 or len(frame.columns) > 100:
                    raise ValueError(f"الورقة {name} تتجاوز حدود نموذج Tracker.")
                header = None
                for r in range(min(12, len(frame))):
                    row = frame.iloc[r]
                    facility_cols = [c for c, v in enumerate(row) if "name of the hospital" in normalize(v) or "name of the phc" in normalize(v)]
                    periods = {c: p for c, v in enumerate(row) if (p := period_for(v, program))}
                    if facility_cols and periods:
                        header = (r, facility_cols[-1], periods)
                        break
                if not header:
                    raise ValueError(f"عناوين الورقة {name} غير مطابقة لنموذج Tracker.")
                sheets.append(name.strip())
                r, name_col, periods = header
                if name_col > 0 and program != "phc_ica":
                    frame.iloc[:, 0] = frame.iloc[:, 0].ffill()
                for row_index in range(r + 1, len(frame)):
                    row = frame.iloc[row_index]
                    facility_name = normalize(row.iloc[name_col])
                    if not facility_name:
                        continue
                    facility = lookup.get((program, facility_name))
                    if program == "hh" and normalize(row.iloc[0]) == "dental centers":
                        alias = {"king abdul aziz hospital": "hh-046", "al noor specialist hospital": "hh-047"}.get(facility_name)
                        facility = ids.get(alias, facility)
                    for column, period in periods.items():
                        try:
                            score = percent(row.iloc[column])
                        except ValueError as exc:
                            raise ValueError(f"{name}، صف {row_index + 1}: {exc}") from None
                        if score is None:
                            continue
                        if score == 0 and skip_zeros:
                            zeros += 1
                            continue
                        if not facility:
                            unmatched.add(f"{name.strip()}: {row.iloc[name_col]}")
                            continue
                        key = (program, facility["id"], period)
                        if key in seen:
                            if seen[key] != score:
                                raise ValueError(f"نتيجتان مختلفتان لنفس المنشأة والفترة: {facility['name']} / {period}. صحح التكرار.")
                            duplicates += 1
                            continue
                        seen[key] = score
                        found.append(dict(program=program, facilityId=facility["id"], facilityName=facility["name"], period=period, score=score, sourceSheet=name.strip(), sourceRow=row_index + 1, group=facility["group"]))
        finally:
            book.close()
        if not found:
            raise ValueError("لم توجد نسب مطابقة قابلة للتحليل. تحقق من عناوين الأوراق والأسماء.")
        return dict(entries=found, unmatched=sorted(unmatched), skippedZeros=zeros, duplicateValues=duplicates, sheets=sheets)


PERIOD_LABELS = {"Q1":"الربع الأول", "Q2":"الربع الثاني", "Q3":"الربع الثالث", "Q4":"الربع الرابع", "H1":"النصف الأول", "H2":"النصف الثاني", "year":"السنة كاملة"}


def selected_periods(scope):
    if scope in ("H1", "H2", "year"):
        return {"H1":["Q1","Q2"], "H2":["Q3","Q4"], "year":["Q1","Q2","Q3","Q4"]}[scope]
    if scope in ("Q1","Q2","Q3","Q4"):
        return [scope]
    raw = scope.removeprefix("compare:").split(",")
    if not scope.startswith("compare:") or not raw or any(q not in ("Q1","Q2","Q3","Q4") for q in raw) or len(set(raw)) != len(raw):
        raise ValueError("اختر فترة صحيحة دون تكرار الأرباع.")
    return sorted(raw)


def scope_label(scope):
    return PERIOD_LABELS.get(scope) or " + ".join(PERIOD_LABELS[q] for q in selected_periods(scope))


class AnalysisAgent:
    """Unweighted complete-facility scores and like-for-like comparisons."""

    def run(self, extracted, catalog, scope="H1", program="all"):
        quarters = selected_periods(scope)
        if program not in ("all", *LABELS):
            raise ValueError("البرنامج المحدد غير معروف.")
        reports = []
        data = pd.DataFrame(extracted["entries"])
        for pid, label in LABELS.items():
            if program != "all" and program != pid:
                continue
            # Half-year measurements can represent complete halves only.
            periods = (["H1","H2"] if quarters == ["Q1","Q2","Q3","Q4"] else
                       ["H1"] if quarters == ["Q1","Q2"] else
                       ["H2"] if quarters == ["Q3","Q4"] else []) if pid == "phc_ica" else quarters
            own = data[data.program == pid]
            lookup = {(e.facilityId,e.period):e.score for e in own.itertuples()}
            rows = []
            for f in [f for f in catalog if f["program"] == pid]:
                values = {}; present = 0
                for period in periods:
                    required = [f"M{m:02}" for m in range((int(period[1])-1)*3+1,int(period[1])*3+1)] if pid == "hh" else [period]
                    source = [lookup.get((f["id"],p)) for p in required]
                    present += sum(v is not None for v in source)
                    values[period] = sum(source)/len(source) if all(v is not None for v in source) else None
                complete = bool(values) and all(v is not None for v in values.values())
                aggregate = sum(values.values())/len(values) if complete else None
                change = values[periods[-1]]-values[periods[0]] if complete and len(periods)>1 else None
                rows.append(dict(name=f.get("arabicName") or f["name"], group=GROUPS.get(f["group"],f["group"]),values=values,aggregate=aggregate,change=change,partial=bool(present and not complete)))
            def avg(values):
                valid=[v for v in values if v is not None]
                return sum(valid)/len(valid) if valid else None
            series=[dict(key=p,label=PERIOD_LABELS[p],mean=avg([r["values"][p] for r in rows]),count=sum(r["values"][p] is not None for r in rows)) for p in periods]
            groups=[]
            for group in dict.fromkeys(r["group"] for r in rows):
                members=[r for r in rows if r["group"]==group]
                groups.append(dict(name=group,values={p:avg([r["values"][p] for r in members]) for p in periods},aggregate=avg([r["aggregate"] for r in members])))
            reports.append(dict(program=pid,label=label,periods=periods,series=series,rows=rows,groups=groups,aggregate=avg([r["aggregate"] for r in rows]),complete=sum(r["aggregate"] is not None for r in rows),total=len(rows),partial=sum(r["partial"] for r in rows),common_count=sum(r["change"] is not None for r in rows),common_change=avg([r["change"] for r in rows])))
        return reports


class SlidePlannerAgent:
    def run(self, analysis, preview, year, scope="H1"):
        label=scope_label(scope)
        plans=[dict(kind="cover",title="تقرير أداء برامج مكافحة العدوى",lines=[label,str(year),"هيئة الصحة العامة · مكتب مكة المكرمة"])]
        plans.append(dict(kind="table",title="الملخص التنفيذي",headers=["البرنامج","متوسط الفترة","مكتملة / بالدليل","إدخال جزئي"],rows=[[r["label"],fmt(r["aggregate"]),f"{r['complete']} / {r['total']}",str(r["partial"])] for r in analysis],note="المتوسط لكل برنامج مستقل، ويشمل المنشآت المكتملة لجميع الفترات المطلوبة فقط. الدليل هو قائمة المنشآت المسجلة بالموقع."))
        plans.append(dict(kind="text",title="منهج الحساب وجودة البيانات",lines=["نظافة اليدين: متوسط 3 أشهر مكتملة لكل ربع. IPCCC وRPP: نتائج ربعية مباشرة.","PHC ICA: نصف سنوي فقط. السنوي يتطلب النصفين؛ لا توزع نتائج النصف على أرباع منفردة.","المتوسطات بسيطة غير موزونة بفرص الرصد. المقارنة المشتركة تتطلب اكتمال جميع الفترات المختارة.",f"أصفار مستبعدة: {preview['skippedZeros']}؛ تكرارات متطابقة: {preview['duplicateValues']}؛ أسماء غير مطابقة: {len(preview['unmatched'])}.","الشرح ناتج عن حسابات محددة دون نموذج لغوي. لا تمثل النتائج أحكامًا سريرية أو حدودًا معيارية."]))
        for r in analysis:
            if not r["periods"]:
                plans.append(dict(kind="text",title=r["label"],lines=["لا تتوفر بيانات ربعية لهذا البرنامج في نموذج Tracker.","اختر نصفًا كاملًا أو السنة كاملة لعرض نتائجه؛ لم يتم توزيع أي نتيجة نصف سنوية على الأرباع."]))
                continue
            note=f"متوسط الفترة: {fmt(r['aggregate'])}. مكتملة: {r['complete']} / {r['total']}؛ جزئية: {r['partial']}."
            if len(r["periods"])>1:
                note+=f" فرق آخر فترة عن الأولى للمراكز المشتركة ({r['common_count']}): {delta(r['common_change'])} نقطة مئوية."
            plans.append(dict(kind="chart",title=r["label"]+" · مقارنة الفترات",report=r,note=note))
            plans.append(dict(kind="table",title=r["label"]+" · تغطية الفترات",headers=["الفترة","المتوسط","مكتملة / بالدليل"],rows=[[v["label"],fmt(v["mean"]),f"{v['count']} / {r['total']}"] for v in r["series"]],note="قد تتغير المنشآت بين الفترات؛ الفرق للمراكز المشتركة في الشريحة السابقة يعزل أثر تغير المجموعة. الفترات الناقصة لا تظهر كأعمدة صفرية."))
            headers=["المنشأة"]+[PERIOD_LABELS[p] for p in r["periods"]]+["متوسط الفترة"]
            for title, rows in [("التصنيفات",r["groups"]),("تفاصيل المنشآت",r["rows"])]:
                plans.append(dict(kind="table",title=r["label"]+" · "+title,headers=["التصنيف" if title=="التصنيفات" else "المنشأة"]+headers[1:],rows=[[v["name"]]+[fmt(v["values"][p]) for p in r["periods"]]+[fmt(v["aggregate"])] for v in rows],note="متوسط الفترة يتطلب اكتمال جميع فتراتها لكل منشأة. غير مكتمل يعني نقص البيانات اللازمة للحساب؛ النتيجة 0% تُحتسب عند اختيار تضمين الأصفار."))
            ranked=sorted([v for v in r["rows"] if v["aggregate"] is not None],key=lambda v:v["aggregate"])
            if ranked:
                plans.append(dict(kind="table",title=r["label"]+" · أولويات المراجعة",headers=["المنشأة","متوسط الفترة","التغير بالنقاط"],rows=[[v["name"],fmt(v["aggregate"]),delta(v["change"])] for v in ranked[:5]],note="أقل النتائج بين المنشآت المكتملة. ترتيب وصفي للمراجعة، دون استنتاج سبب أو مخالفة معيار غير محدد."))
        plans.append(dict(kind="text",title="المتابعة قبل اعتماد التقرير",lines=["استكمال الفترات الناقصة ومراجعة الأسماء غير المطابقة قبل اعتماد النتائج.","مناقشة التغيرات مع مسؤولي البرامج وتوثيق الأسباب المؤكدة والإجراء والموعد.","يمكن تحرير نصوص العرض وجداوله ورسومه في PowerPoint قبل الطباعة أو التصدير إلى PDF.","مراجعة: ____________________     اعتماد: ____________________"]))
        return plans


class PptxGeneratorAgent(BaseRenderer):
    NAVY = RGBColor.from_string("114C50")
    TEAL = RGBColor.from_string("0C938E")
    PALE = RGBColor.from_string("EDF7F5")

    def run(self, plans, year, logo=None, scope="H1"):
        cfg = ReportConfig(year=year, sheets=(SheetSpec(name="Tracker"),), title="تقرير برامج وقاية", organization="هيئة الصحة العامة · مكتب مكة المكرمة")
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        deck.core_properties.title = cfg.title
        deck.core_properties.author = cfg.organization
        for plan in plans:
            pages = [plan.get("rows", [])[i:i+7] for i in range(0, len(plan.get("rows", [])), 7)] or [[]]
            for page_index, rows in enumerate(pages):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, Inches(.12))
                band.fill.solid(); band.fill.fore_color.rgb = self.TEAL; band.line.fill.background()
                self._text(slide, plan["title"] + (" · تابع" if page_index else ""), 2.2, .35, 10.45, .9, cfg, size=26, color=self.NAVY, bold=True)
                if logo:
                    slide.shapes.add_picture(io.BytesIO(logo), Inches(.6), Inches(.25), width=Inches(1.35))
                self._text(slide, f"{year}  |  {scope_label(scope)}  |  {len(deck.slides)}", .6, 7.06, 12.1, .25, cfg, size=10, color=self.TEAL)
                if plan["kind"] == "cover":
                    for i, line in enumerate(plan["lines"]):
                        self._text(slide,line,.9,2.1+i*1.08,11.55,.88,cfg,size=32 if i==0 else 24,color=self.NAVY,bold=i==0)
                elif plan["kind"] == "text":
                    for i, line in enumerate(plan["lines"]):
                        self._text(slide,line,.8,1.65+i*.98,11.7,.88,cfg,size=18)
                elif plan["kind"] == "table":
                    self._report_table(slide,plan["headers"],rows,cfg)
                    self._text(slide,plan["note"],.8,6.22,11.7,.65,cfg,size=13,color=self.NAVY)
                else:
                    r=plan["report"]
                    values=[(v["label"],v["mean"]) for v in r["series"]]
                    available=[(k,v) for k,v in values if v is not None]
                    if available:
                        chart_data=CategoryChartData();chart_data.categories=[k for k,v in available];chart_data.add_series("المتوسط %",[v for k,v in available])
                        chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(11.2), Inches(4.35), chart_data).chart
                        chart.has_legend=False;chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=100
                        chart.value_axis.tick_labels.font.size=Pt(13);chart.category_axis.tick_labels.font.size=Pt(16)
                        chart.plots[0].has_data_labels=True;chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
                        chart.plots[0].data_labels.number_format='0.0"%"';chart.plots[0].data_labels.font.size=Pt(16)
                        chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=self.TEAL
                    else:
                        self._text(slide,"لا توجد نتائج مكتملة للفترة المحددة",1,3,11.2,1,cfg,size=26)
                    self._text(slide,plan["note"],.8,6.22,11.7,.65,cfg,size=14,color=self.NAVY)
        output=io.BytesIO();deck.save(output);return output.getvalue()

    def _report_table(self,slide,headers,rows,cfg):
        table=slide.shapes.add_table(len(rows)+1,len(headers), Inches(.75), Inches(1.5), Inches(11.8), Inches(.55*(len(rows)+1))).table
        for col in table.columns: col.width=Inches(7.0/(len(headers)-1))
        table.columns[len(headers)-1].width=Inches(4.8)
        for ri, values in enumerate([headers]+rows):
            for ci, value in enumerate(reversed(values)):
                cell=table.cell(ri,ci);cell.text=str(value);cell.fill.solid();cell.fill.fore_color.rgb=self.NAVY if ri==0 else self.PALE if ri%2 else RGBColor(255,255,255)
                cell.margin_top=Inches(.08);cell.margin_bottom=Inches(.05)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph._p.get_or_add_pPr().set("rtl","1")
                    for run in paragraph.runs:
                        run.font.name="Arial";run.font.size=Pt(13 if ci!=len(headers)-1 else 14);run.font.bold=ri==0
                        run.font.color.rgb=RGBColor(255,255,255) if ri==0 else self.INK


@dataclass
class GeneratedReport:
    pptx_bytes: bytes
    preview: dict
    analysis: list


class InfectionControlSupervisor:
    """Coordinate four specialized agents and validate their final artifact."""

    def run(self, source: bytes, catalog: list[dict], year=2026, logo=None, skip_zeros=True, progress: Callable[[str],None] | None=None, scope="H1", program="all"):
        if not 2020 <= int(year) <= 2100:
            raise ValueError("سنة التقرير غير صحيحة.")
        status=progress or (lambda message: None)
        status("مطابقة المنشآت والفترات")
        preview=DataExtractorAgent().run(source,catalog,skip_zeros)
        status("تحليل الفترات المختارة والتغطية")
        analysis=AnalysisAgent().run(preview,catalog,scope,program)
        status("تنظيم المقارنات والشروح")
        plan=SlidePlannerAgent().run(analysis,preview,year,scope)
        status("تصميم العرض والجداول والرسوم")
        output=PptxGeneratorAgent().run(plan,year,logo,scope)
        status("مراجعة العرض النهائي")
        with zipfile.ZipFile(io.BytesIO(output)) as archive:
            if archive.testzip() or "ppt/presentation.xml" not in archive.namelist():
                raise ValueError("تعذر التحقق من ملف العرض.")
        return GeneratedReport(output,preview,analysis)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel",type=Path);parser.add_argument("--output",type=Path,default=Path("report.pptx"));parser.add_argument("--year",type=int,default=2026)
    parser.add_argument("--include-zero",action="store_true")
    parser.add_argument("--scope",default="H1",help="Q1..Q4, H1, H2, year or compare:Q1,Q3")
    parser.add_argument("--program",default="all",choices=["all",*LABELS])
    args=parser.parse_args();root=Path(__file__).parent
    result=InfectionControlSupervisor().run(args.excel.read_bytes(),json.loads((root/"lib/facilities.json").read_text(encoding="utf-8")),args.year,(root/"public/wiqaya-logo.png").read_bytes(),not args.include_zero,scope=args.scope,program=args.program)
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_bytes(result.pptx_bytes)
    print(f"Created {args.output}; imported measurements={len(result.preview['entries'])}; deterministic mode")


if __name__=="__main__": main()
